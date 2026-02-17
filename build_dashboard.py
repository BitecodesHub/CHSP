#!/usr/bin/env python3
"""Generate enhanced CHSP dashboard from Report.xlsx"""
import openpyxl, json

wb = openpyxl.load_workbook('Report.xlsx', data_only=True)
ws = wb['data']
headers = [cell.value for cell in ws[1]]
data = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    rec = {}
    for j, h in enumerate(headers):
        v = row[j]
        if v is None: rec[h] = ""
        elif isinstance(v, float): rec[h] = int(v) if v == int(v) else round(v, 2)
        else: rec[h] = v
    data.append(rec)

json_str = json.dumps(data)

# Read template parts
with open('tmpl_head.html') as f: head = f.read()
with open('tmpl_body.html') as f: body = f.read()
with open('tmpl_script.js') as f: script = f.read()

html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>CHSP Allied Health — Advanced Dashboard</title>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800;900&display=swap" rel="stylesheet">
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js"></script>
{head}
</head>
<body>
{body}
<script>
const DATA = {json_str};
{script}
</script>
</body>
</html>"""

with open('index.html', 'w') as f:
    f.write(html)
print("Dashboard generated successfully!")
